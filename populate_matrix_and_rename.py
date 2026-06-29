import os
import re
import glob
import openpyxl
from openpyxl.styles import Alignment

EXTRACTED_DIR = r"d:\Research\extracted_texts"
PAPERS_DIR = r"d:\Research\papers"
EXCEL_FILE = r"d:\Research\Literature_Review_Matrix.xlsx"

def clean_filename(title):
    # Keep only alphanumeric and spaces, then replace spaces with underscores
    clean = re.sub(r'[^a-zA-Z0-9\s]', '', title)
    clean = clean.replace('\n', ' ').strip()
    return "_".join(clean.split())

def extract_paper_data():
    paper_data = {}
    batch_files = glob.glob(os.path.join(EXTRACTED_DIR, "batch_*.txt"))
    
    for batch_file in batch_files:
        with open(batch_file, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
            
        papers = re.split(r"--- START OF PAPER: (.*?) ---", content)
        for i in range(1, len(papers), 2):
            filename = papers[i].strip()
            text = papers[i+1].strip()
            
            lines = text.split('\n')
            title_lines = []
            for line in lines[:20]:
                line = line.strip()
                if not line:
                    continue
                if any(x in line.upper() for x in ["THIS ARTICLE HAS BEEN", "IEEE", "COPYRIGHT", "RECEIVED", "ACCEPTED", "DOI"]):
                    continue
                if "Abstract" in line or "ABSTRACT" in line:
                    break
                title_lines.append(line)
            title = " ".join(title_lines[:3]).strip()
            if not title or len(title) < 5:
                title = filename.replace(".pdf", "")
                
            methodology = "Not extracted"
            method_match = re.search(r'\n(III\.|[IVX]+\.|[0-9]+\.)?\s*(Proposed Method|Methodology|Methods|Materials and Methods)(.*?)\n[IVX0-9]+\.', text, re.IGNORECASE | re.DOTALL)
            if method_match:
                methodology = method_match.group(3)[:500].strip().replace('\n', ' ') + "..."
            else:
                abs_match = re.search(r'(Abstract|ABSTRACT)(.*?)(Introduction|1\.)', text, re.IGNORECASE | re.DOTALL)
                if abs_match:
                    methodology = "From Abstract: " + abs_match.group(2)[:500].strip().replace('\n', ' ') + "..."
            
            gaps = "Not extracted"
            gaps_match = re.search(r'\n(VI\.|VII\.|[IVX]+\.|[0-9]+\.)?\s*(Conclusion|Discussion|Future Work)(.*?)(References|Acknowledgment)', text, re.IGNORECASE | re.DOTALL)
            if gaps_match:
                content_gaps = gaps_match.group(3)
                future_work = re.search(r'(future work|future research|in the future)(.*?)\.', content_gaps, re.IGNORECASE | re.DOTALL)
                if future_work:
                    gaps = future_work.group(0).strip().replace('\n', ' ')
                else:
                    gaps = content_gaps[:500].strip().replace('\n', ' ') + "..."
            
            paper_data[filename] = {
                "title": title,
                "methodology": methodology,
                "gaps": gaps
            }
    return paper_data

def rename_pdfs(paper_data):
    rename_mapping = {}
    for old_filename, data in paper_data.items():
        old_path = os.path.join(PAPERS_DIR, old_filename)
        if os.path.exists(old_path):
            new_name = clean_filename(data["title"])[:100] + ".pdf"
            new_path = os.path.join(PAPERS_DIR, new_name)
            if old_path != new_path:
                try:
                    os.rename(old_path, new_path)
                    print(f"Renamed: {old_filename} -> {new_name}")
                    rename_mapping[old_filename] = new_name
                except Exception as e:
                    print(f"Failed to rename {old_filename}: {e}")
                    rename_mapping[old_filename] = old_filename
            else:
                rename_mapping[old_filename] = old_filename
        else:
            rename_mapping[old_filename] = old_filename
    return rename_mapping

def update_excel(paper_data):
    wb = openpyxl.load_workbook(EXCEL_FILE)
    
    ws_matrix = wb['📚 Literature Matrix']
    row_idx = 5
    for filename, data in paper_data.items():
        paper_id = f"P{row_idx-4:03d}"
        ws_matrix.cell(row=row_idx, column=1, value=paper_id)
        ws_matrix.cell(row=row_idx, column=2, value=data["title"])
        row_idx += 1
        
    ws_gaps = wb['🔍 Gap Tracker']
    row_idx = 4
    for filename, data in paper_data.items():
        paper_id = f"P{row_idx-3:03d}"
        ws_gaps.cell(row=row_idx, column=1, value=paper_id)
        ws_gaps.cell(row=row_idx, column=2, value=data["title"])
        c = ws_gaps.cell(row=row_idx, column=4, value=data["gaps"])
        c.alignment = Alignment(wrap_text=True)
        row_idx += 1

    ws_method = wb['⚙️ Methodology Comparison']
    row_idx = 4
    for filename, data in paper_data.items():
        paper_id = f"P{row_idx-3:03d}"
        ws_method.cell(row=row_idx, column=1, value=paper_id)
        ws_method.cell(row=row_idx, column=2, value=data["title"])
        c = ws_method.cell(row=row_idx, column=3, value=data["methodology"])
        c.alignment = Alignment(wrap_text=True)
        row_idx += 1

    wb.save(EXCEL_FILE)
    print("Excel updated.")

if __name__ == "__main__":
    data = extract_paper_data()
    rename_pdfs(data)
    update_excel(data)
    print("Done!")
